#!/usr/bin/env python3
"""
split_wht_tabs.py — Split WHT certificate copies into separate tabs

Each sheet in WHT_xxxx.xlsx has 4 certificate copies side-by-side.
Copies 2–4 use formulas referencing copy 1's cells, so we clone the
full sheet and set the print area (rather than extracting columns),
keeping all formula links intact.

Produces two output workbooks:
  WHT_xxxx_all4.xlsx  — all 4 copies as separate tabs per person
  WHT_xxxx_fast.xlsx  — copies 3 & 4 only per person

Usage:
    python split_wht_tabs.py output/WHT/WHT_6905.xlsx
"""

import os
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# Each copy's *body* column range (the gutter/spacer columns S=19 and BD=56
# that sit between side-by-side copies are EXCLUDED so all four copies share
# the exact same printable width -> identical size when printed individually).
COPIES = [
    (1,  18, "c1"),   # A-R   (col S is a blank spacer, dropped)
    (20, 37, "c2"),   # T-AK
    (38, 55, "c3"),   # AL-BC (col BD is a blank spacer, dropped)
    (57, 74, "c4"),   # BE-BV
]


def build_workbook(xlsx_path, include_suffixes):
    wb = load_workbook(xlsx_path)
    original_names = list(wb.sheetnames)

    for name in original_names:
        src = wb[name]
        for col_start, col_end, suffix in COPIES:
            if suffix not in include_suffixes:
                continue
            clone = wb.copy_worksheet(src)
            clone.title = f"{name[:27]}-{suffix}"

            # Print area = this copy's body columns only (no gutter)
            clone.print_area = (
                f"{get_column_letter(col_start)}1:"
                f"{get_column_letter(col_end)}{src.max_row}"
            )

            # Pin identical page setup on EVERY tab so each cert prints at the
            # same size regardless of the print dialog. Fit-to-one-page + the
            # now-uniform print width => every cert fills one A4 identically.
            clone.page_setup.orientation = "landscape"
            clone.page_setup.paperSize = 9          # A4
            clone.page_setup.fitToWidth = 1
            clone.page_setup.fitToHeight = 1
            clone.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            clone.print_options.horizontalCentered = True
            clone.print_options.verticalCentered = True
            clone.page_margins.left = clone.page_margins.right = 0.22
            clone.page_margins.top = clone.page_margins.bottom = 0.11
            clone.page_margins.header = clone.page_margins.footer = 0.0
            clone.sheet_view.showGridLines = False

            # Hide columns outside this copy so each tab looks like one cert
            for col_idx in range(1, src.max_column + 1):
                ltr = get_column_letter(col_idx)
                if col_idx < col_start or col_idx > col_end:
                    clone.column_dimensions[ltr].hidden = True

    # Remove the originals
    for name in original_names:
        del wb[name]

    return wb


def main():
    if len(sys.argv) < 2:
        print("Usage: python split_wht_tabs.py <WHT_xxxx.xlsx>")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    if not os.path.exists(xlsx_path):
        print(f"File not found: {xlsx_path}")
        sys.exit(1)

    base = os.path.splitext(xlsx_path)[0]

    out_all4 = base + "_all4.xlsx"
    build_workbook(xlsx_path, {"c1", "c2", "c3", "c4"}).save(out_all4)
    print(f"All 4 copies  -> {out_all4}")

    out_fast = base + "_fast.xlsx"
    build_workbook(xlsx_path, {"c3", "c4"}).save(out_fast)
    print(f"Fast (3 & 4)  -> {out_fast}")


if __name__ == "__main__":
    main()
