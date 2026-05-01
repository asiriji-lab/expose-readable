"""
Schedule Excel Exporter

Converts schedule JSON + entity_meta into a 3-sheet Excel workbook:
    Teacher, Student, Room

Layout per sheet:
  Row 1:  global header — col A label ("คาบ"), cols B+ = period labels
  Row 2:  global times  — col A label ("เวลา"), cols B+ = period times
  Then for each entity a 16-row block followed by 1 blank separator:
    Row  1: entity label (col A) + period times repeated (cols B+)
    Rows 2-4:  จันทร์  (col A merged, 3 rows)
    Rows 5-7:  อังคาร
    Rows 8-10: พุธ
    Rows 11-13:พฤหัส
    Rows 14-16:ศุกร์

  Each (day × period) slot occupies 3 rows:
    top  row: subject_id
    mid  row: teacher  (student view) | class  (teacher view) | teacher (room view)
    bot  row: room     (student view) | room   (teacher view) | class   (room view)
"""

import io
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font

WEEKDAYS       = ['MON', 'TUE', 'WED', 'THU', 'FRI']
WEEKDAYS_FULL  = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
DAY_THAI       = ['จันทร์', 'อังคาร', 'พุธ', 'พฤหัส', 'ศุกร์']
_ROWS_PER_DAY   = 3
_DAYS_PER_BLOCK = 5
_ENTITY_ROWS    = 1 + _DAYS_PER_BLOCK * _ROWS_PER_DAY  # 16
_SEPARATOR_ROWS = 1


# ---------------------------------------------------------------------------
# Format normaliser  (frontend FullDataset → backend list-of-entity-dicts)
# ---------------------------------------------------------------------------

def _normalize_schedule(schedule_json: Dict, entity_meta: Optional[Dict]) -> Dict:
    """
    The frontend saves in FullDataset format:
        teachers / classes / rooms  →  dict keyed by entity code
        each value                  →  { day: { slotIndex: ScheduleItem } }

    The exporter expects:
        teachers / students / rooms →  list of { id, name, rows: [...] }

    Detect the frontend format (teachers is a dict, not a list) and convert.
    Backend format is returned unchanged.
    """
    teachers_raw = schedule_json.get('teachers', [])
    if isinstance(teachers_raw, list):
        return schedule_json  # already backend format

    # ── frontend format detected ──────────────────────────────────────────────
    classes_raw = schedule_json.get('classes', {})
    rooms_raw   = schedule_json.get('rooms', {})
    teacher_meta = (entity_meta or {}).get('teacher_meta', {})

    # Collect all slot indices present in the data so we know the period range.
    # Frontend FullDataset: entity_map[entityCode][dayFullName][slotInt] = item
    # We must go 3 levels deep: entity → day → slot.
    all_slots: set = set()
    for entity_map in (teachers_raw, classes_raw, rooms_raw):
        for entity_schedule in entity_map.values():
            if not isinstance(entity_schedule, dict):
                continue
            for day_data in entity_schedule.values():
                if not isinstance(day_data, dict):
                    continue
                all_slots.update(str(k) for k in day_data.keys())
    period_labels = sorted(all_slots, key=lambda x: int(x) if x.isdigit() else 0)

    # Build / patch config.columns so _parse_periods works correctly.
    config = dict(schedule_json.get('config') or {})
    if not config.get('columns'):
        columns: List[Dict] = [{'key': 'day', 'label': 'Day', 'type': 'text'}]
        for lbl in period_labels:
            columns.append({'key': lbl, 'label': lbl, 'time': ''})
        config['columns'] = columns

    def _to_cell(item: Dict, view: str) -> Optional[Dict]:
        sid = item.get('subjectCode') or ''
        if not sid:
            return None
        if view == 'teacher':
            return {'subject_id': sid, 'class': item.get('classCode') or '', 'room': item.get('room') or ''}
        if view == 'student':
            return {'subject_id': sid, 'teacher': item.get('teacherName') or item.get('teacher') or '', 'room': item.get('room') or ''}
        # room
        return {'subject_id': sid, 'teacher': item.get('teacherName') or item.get('teacher') or '', 'class': item.get('classCode') or ''}

    def _build_rows(entity_schedule: Dict, view: str) -> List[Dict]:
        rows = []
        # Frontend stores days as full English names; try full name first, abbrev as fallback.
        for day_abbr, day_full in zip(WEEKDAYS, WEEKDAYS_FULL):
            day_data = entity_schedule.get(day_full) or entity_schedule.get(day_abbr) or {}
            columns: List[Dict] = []
            for lbl in period_labels:
                # Frontend stores slot keys as ints; try int first, then string.
                item = day_data.get(int(lbl)) if lbl.isdigit() else None
                if item is None:
                    item = day_data.get(lbl)
                columns.append({lbl: _to_cell(item, view) if item else None})
            rows.append({'day': day_abbr, 'columns': columns})
        return rows

    teachers = [
        {
            'id': code,
            'name': (teacher_meta.get(code) or {}).get('name', code),
            'rows': _build_rows(sched, 'teacher'),
        }
        for code, sched in teachers_raw.items()
    ]
    students = [
        {'id': code, 'name': code, 'rows': _build_rows(sched, 'student')}
        for code, sched in classes_raw.items()
    ]
    rooms = [
        {'id': code, 'name': code, 'rows': _build_rows(sched, 'room')}
        for code, sched in rooms_raw.items()
    ]

    return {
        'config':   config,
        'teachers': teachers,
        'students': students,
        'rooms':    rooms,
    }


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def build_schedule_excel(
    schedule_json: Dict,
    entity_meta: Optional[Dict] = None,
) -> bytes:
    """Return Excel workbook bytes for the given schedule."""
    schedule_json = _normalize_schedule(schedule_json, entity_meta)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    config       = schedule_json.get('config', {})
    period_labels, period_times = _parse_periods(config)
    class_meta   = (entity_meta or {}).get('class_meta', {})

    _build_teacher_sheet(wb, schedule_json, period_labels, period_times)
    _build_student_sheet(wb, schedule_json, period_labels, period_times, class_meta)
    _build_room_sheet(wb, schedule_json, period_labels, period_times)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---------------------------------------------------------------------------
# Period helpers
# ---------------------------------------------------------------------------

def _parse_periods(config: Dict) -> Tuple[List[str], Dict[str, str]]:
    """Return (ordered period labels, label→time map) from config.columns."""
    labels: List[str] = []
    times:  Dict[str, str] = {}
    for col in config.get('columns', []):
        if col.get('key') == 'day':
            continue
        label = col['key']
        labels.append(label)
        times[label] = col.get('time', '')
    return labels, times


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_teacher_sheet(
    wb,
    schedule_json: Dict,
    period_labels: List[str],
    period_times: Dict[str, str],
) -> None:
    ws = wb.create_sheet('Teacher')
    _write_global_header(ws, period_labels, period_times)
    row = 3
    for teacher in schedule_json.get('teachers', []):
        label = teacher.get('name') or teacher.get('id', '')
        row = _write_entity_block(ws, teacher, 'teacher', period_labels, period_times, row, label)


def _build_student_sheet(
    wb,
    schedule_json: Dict,
    period_labels: List[str],
    period_times: Dict[str, str],
    class_meta: Dict,
) -> None:
    ws = wb.create_sheet('Student')
    _write_global_header(ws, period_labels, period_times)
    row = 3
    for student in schedule_json.get('students', []):
        class_id = student.get('id', '')
        homeroom = (class_meta.get(class_id) or {}).get('defaultRoom', '')
        label = f"{class_id} {homeroom}".strip() if homeroom else class_id
        row = _write_entity_block(ws, student, 'student', period_labels, period_times, row, label)


def _build_room_sheet(
    wb,
    schedule_json: Dict,
    period_labels: List[str],
    period_times: Dict[str, str],
) -> None:
    ws = wb.create_sheet('Room')
    _write_global_header(ws, period_labels, period_times)
    row = 3
    for room in schedule_json.get('rooms', []):
        label = room.get('name') or room.get('id', '')
        row = _write_entity_block(ws, room, 'room', period_labels, period_times, row, label)


# ---------------------------------------------------------------------------
# Row writers
# ---------------------------------------------------------------------------

def _write_global_header(
    ws,
    period_labels: List[str],
    period_times: Dict[str, str],
) -> None:
    """Write rows 1 and 2: period numbers and period times."""
    ws.cell(row=1, column=1, value='คาบ').font = Font(bold=True)
    ws.cell(row=2, column=1, value='เวลา').font = Font(bold=True)
    for i, label in enumerate(period_labels):
        col = i + 2
        ws.cell(row=1, column=col, value=label).font = Font(bold=True)
        ws.cell(row=2, column=col, value=period_times.get(label, '')).font = Font(bold=True)


def _write_entity_block(
    ws,
    entity: Dict,
    view: str,
    period_labels: List[str],
    period_times: Dict[str, str],
    start_row: int,
    entity_label: str,
) -> int:
    """
    Write 16-row entity block + 1 blank separator.
    Returns the next available start_row.
    """
    # Header row: entity label + period times
    hdr = ws.cell(row=start_row, column=1, value=entity_label)
    hdr.font = Font(bold=True)
    hdr.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for i, label in enumerate(period_labels):
        col = i + 2
        ws.cell(row=start_row, column=col, value=period_times.get(label, ''))

    # Build day→{period_label: cell_dict} index
    rows_by_day: Dict[str, Dict[str, Any]] = {}
    for row_data in entity.get('rows', []):
        day = row_data['day']
        cells: Dict[str, Any] = {}
        for col_dict in row_data.get('columns', []):
            for k, v in col_dict.items():
                cells[k] = v
        rows_by_day[day] = cells

    # Write 5 days
    for d_idx, (day, day_thai) in enumerate(zip(WEEKDAYS, DAY_THAI)):
        top_row = start_row + 1 + d_idx * _ROWS_PER_DAY
        mid_row = top_row + 1
        bot_row = top_row + 2

        # Merge col A for day label
        ws.merge_cells(
            start_row=top_row, start_column=1,
            end_row=bot_row,   end_column=1,
        )
        day_cell = ws.cell(row=top_row, column=1, value=day_thai)
        day_cell.font      = Font(bold=True)
        day_cell.alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True
        )

        day_cells = rows_by_day.get(day, {})
        for i, label in enumerate(period_labels):
            col  = i + 2
            cell = day_cells.get(label)
            top, mid, bot = _slot_values(cell, view)
            if top:
                ws.cell(row=top_row, column=col, value=top)
            if mid:
                ws.cell(row=mid_row, column=col, value=mid)
            if bot:
                ws.cell(row=bot_row, column=col, value=bot)

    return start_row + _ENTITY_ROWS + _SEPARATOR_ROWS


# ---------------------------------------------------------------------------
# Cell value extractor
# ---------------------------------------------------------------------------

def _slot_values(
    cell: Optional[Dict],
    view: str,
) -> Tuple[Optional[str], Optional[str], Optional[str]]:
    """
    Return (top, mid, bot) for a schedule cell.

    student : subject_id | teacher | room
    teacher : subject_id | class   | room
    room    : subject_id | teacher | class
    """
    if not cell:
        return None, None, None
    sid = cell.get('subject_id') or ''
    if not sid:
        return None, None, None
    if view == 'student':
        return sid, cell.get('teacher') or '', cell.get('room') or ''
    if view == 'teacher':
        return sid, cell.get('class') or '', cell.get('room') or ''
    # room
    return sid, cell.get('teacher') or '', cell.get('class') or ''
